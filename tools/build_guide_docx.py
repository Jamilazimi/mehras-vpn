"""
ساخت فایل راهنمای «وی پی ان شخصی مهراس» به صورت DOCX (راست‌به‌چپ، فونت وزیرمتن).

اجرا با پایتون بسته‌ی محیط:
    python tools/build_guide_docx.py
سپس با LibreOffice به PDF تبدیل می‌شود (tools/export_guide_pdf.ps1).
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "راهنمای-مهراس.docx"

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

FONT = "Vazirmatn"
NAVY = RGBColor(0x10, 0x1B, 0x2D)
ACCENT = RGBColor(0x1F, 0x6F, 0xE0)
GRAY = RGBColor(0x55, 0x62, 0x74)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BODY_PT = 10.5

FA_DIGITS = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")


def fa(text) -> str:
    return str(text).translate(FA_DIGITS)


# ---------------------------------------------------------------------------
# کمک‌تابع‌های XML
# ---------------------------------------------------------------------------
PPR_AFTER_PBDR = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
    "w:pPrChange",
)
PPR_AFTER_BIDI = (
    "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
    "w:pPrChange",
)
TBLPR_AFTER_BORDERS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
                       "w:tblCaption", "w:tblDescription")


def _set(el, tag, **attrs):
    child = OxmlElement(tag)
    for key, value in attrs.items():
        child.set(qn(f"w:{key}"), str(value))
    el.append(child)
    return child


def mark_rtl_paragraph(p):
    """افزودن <w:bidi/> در جای درست (ترتیب المان‌های pPr مهم است)."""
    p_pr = p._p.get_or_add_pPr()
    if p_pr.find(qn("w:bidi")) is not None:
        return p_pr
    p_pr.insert_element_before(OxmlElement("w:bidi"), *PPR_AFTER_BIDI)
    return p_pr


def shade_cell(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.insert_element_before(
        shd, "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
        "w:vAlign", "w:hideMark", "w:headers", "w:cellIns", "w:cellDel",
        "w:cellMerge", "w:tcPrChange")


def cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("bottom", bottom),
                        ("start", left), ("end", right)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tc_pr.insert_element_before(
        mar, "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark",
        "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange")


def table_borders(table, color="D6E0F0", size=6):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.insert_element_before(borders, *TBLPR_AFTER_BORDERS)


def paragraph_bottom_border(par, color="1F6FE0", size=8, space=3):
    p_pr = par._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.insert_element_before(p_bdr, *PPR_AFTER_PBDR)


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        _set(tr_pr, "w:cantSplit")


def add_page_number_footer(section):
    footer = section.footer
    par = footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run("صفحه ")
    _style_run(run, size=8.5, color=GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    par._p.append(fld)
    run2 = par.add_run(" از ")
    _style_run(run2, size=8.5, color=GRAY)
    fld2 = OxmlElement("w:fldSimple")
    fld2.set(qn("w:instr"), "NUMPAGES")
    par._p.append(fld2)


def add_header(section, text):
    header = section.header
    par = header.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = par.add_run(text)
    _style_run(run, size=8.5, color=GRAY)
    paragraph_bottom_border(par, color="E2E8F4", size=4, space=2)


# ---------------------------------------------------------------------------
# کمک‌تابع‌های متن
# ---------------------------------------------------------------------------
def _style_run(run, size=BODY_PT, bold=False, color=None, italic=False, rtl=True):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(attr), FONT)
    # اندازه/برجستگی خط پیچیده (فارسی) — Word برای متن فارسی از szCs/bCs استفاده می‌کند
    sz = r_pr.find(qn("w:sz"))
    if sz is not None and r_pr.find(qn("w:szCs")) is None:
        sz_cs = OxmlElement("w:szCs")
        sz_cs.set(qn("w:val"), str(int(size * 2)))
        sz.addnext(sz_cs)
    b = r_pr.find(qn("w:b"))
    if b is not None and r_pr.find(qn("w:bCs")) is None:
        b_cs = OxmlElement("w:bCs")
        b_cs.set(qn("w:val"), "1" if bold else "0")
        b.addnext(b_cs)
    run.font.rtl = rtl
    return run


def par(doc, text="", size=BODY_PT, bold=False, color=NAVY, align="justify",
        space_before=0, space_after=6, line=1.3, rtl=True, indent=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if indent:
        pf.right_indent = Cm(indent)
    pf.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
                    "right": WD_ALIGN_PARAGRAPH.RIGHT,
                    "center": WD_ALIGN_PARAGRAPH.CENTER,
                    "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    if rtl:
        mark_rtl_paragraph(p)
    if text:
        _style_run(p.add_run(text), size=size, bold=bold, color=color, rtl=rtl)
    return p


def rich_par(doc, parts, align="justify", space_after=6, size=BODY_PT):
    """پاراگراف با چند بخش (متن عادی + برجسته + لاتین)."""
    p = par(doc, "", size=size, align=align, space_after=space_after)
    for part in parts:
        if isinstance(part, str):
            _style_run(p.add_run(part), size=size, color=NAVY)
        else:
            text, opts = part
            _style_run(p.add_run(text), size=opts.get("size", size),
                       bold=opts.get("bold", False),
                       color=opts.get("color", NAVY),
                       rtl=opts.get("rtl", True))
    return p


def h1(doc, text, number=None):
    # برای پاراگراف راست‌به‌چپ، «both» مطمئن‌ترین حالت است (تک‌خطی = شروع از راست)
    p = par(doc, "", size=15, space_before=14, space_after=8, align="justify")
    if number:
        _style_run(p.add_run(f"{fa(number)}. "), size=15, bold=True, color=ACCENT)
    _style_run(p.add_run(text), size=15, bold=True, color=NAVY)
    paragraph_bottom_border(p, color="1F6FE0", size=8, space=4)
    return p


def h2(doc, text):
    return par(doc, text, size=11.5, bold=True, color=ACCENT,
               space_before=8, space_after=3, align="justify")


def bullet(doc, text, marker="•", bold_prefix=None):
    p = par(doc, "", align="right", space_after=2, indent=0.2)
    _style_run(p.add_run(f"{marker}  "), size=BODY_PT, color=ACCENT, bold=True)
    if bold_prefix:
        _style_run(p.add_run(bold_prefix), size=BODY_PT, bold=True, color=NAVY)
    _style_run(p.add_run(text), size=BODY_PT, color=NAVY)
    return p


def callout(doc, title, text, fill="EEF4FF", border="1F6FE0"):
    table_ = doc.add_table(rows=1, cols=1)
    table_.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table_.cell(0, 0)
    shade_cell(cell, fill)
    cell_margins(cell, 140, 140, 180, 180)
    tbl_pr = table_._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "single")
    start.set(qn("w:sz"), "24")
    start.set(qn("w:space"), "0")
    start.set(qn("w:color"), border)
    borders.append(start)
    tbl_pr.insert_element_before(borders, *TBLPR_AFTER_BORDERS)
    keep_row_together(table_.rows[0])

    p = cell.paragraphs[0]
    mark_rtl_paragraph(p)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _style_run(p.add_run(title), size=10.5, bold=True, color=ACCENT)
    p2 = cell.add_paragraph()
    mark_rtl_paragraph(p2)
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.3
    p2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _style_run(p2.add_run(text), size=10, color=NAVY)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def table(doc, headers, rows, widths=None, header_fill="E8EFFB",
          font_size=9.5, ltr_cols=()):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(tbl)
    hdr = tbl.rows[0]
    keep_row_together(hdr)
    # سرصفحه در صفحات بعدی هم تکرار شود
    tr_pr = hdr._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        _set(tr_pr, "w:tblHeader")
    for i, title in enumerate(headers):
        cell = hdr.cells[i]
        shade_cell(cell, header_fill)
        cell_margins(cell)
        p = cell.paragraphs[0]
        mark_rtl_paragraph(p)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        _style_run(p.add_run(title), size=font_size, bold=True, color=NAVY)
    for values in rows:
        row = tbl.add_row()
        keep_row_together(row)
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell_margins(cell)
            p = cell.paragraphs[0]
            ltr = i in ltr_cols
            if not ltr:
                mark_rtl_paragraph(p)
            p.paragraph_format.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if ltr else WD_ALIGN_PARAGRAPH.RIGHT)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.2
            _style_run(p.add_run(str(value)), size=font_size, color=NAVY, rtl=not ltr)
    if widths:
        for row in tbl.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return tbl


# ---------------------------------------------------------------------------
# ساخت سند
# ---------------------------------------------------------------------------
def build():
    doc = Document()

    # استایل پیش‌فرض
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(BODY_PT)
    style.element.rPr.rFonts.set(qn("w:cs"), FONT)

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    # جهت راست‌به‌چپ برای کل بخش
    sect_pr = section._sectPr
    _set(sect_pr, "w:bidi")
    add_header(section, "وی پی ان شخصی مهراس — راهنمای استفاده")
    add_page_number_footer(section)

    # ---------------- جلد ----------------
    par(doc, "", space_after=40)
    title = par(doc, "", align="center", space_after=4)
    _style_run(title.add_run("وی پی ان شخصی مهراس"), size=30, bold=True, color=NAVY)
    sub = par(doc, "", align="center", space_after=2)
    _style_run(sub.add_run("Mehras VPN"), size=12, color=ACCENT)
    line = par(doc, "", align="center", space_after=18)
    _style_run(line.add_run("راهنمای کوتاه استفاده — نسخه ۲٫۰"), size=13,
               color=GRAY)

    intro = ("مهراس یک کلاینت شخصی برای عبور از فیلترینگ است. برنامه خودش "
             "گره‌های تازه را از منابع معتبر می‌گیرد، با اتصال واقعی تست می‌کند، "
             "بهترین را انتخاب می‌کند و اگر اتصال ضعیف شد خودکار جایگزین "
             "می‌کند. لازم نیست هیچ کانفیگی را دستی وارد کنید.")
    par(doc, intro, size=11, space_after=18)

    table(doc, ["موضوع", "توضیح"], [
        ["نام برنامه", "وی پی ان شخصی مهراس (نسخه ۲٫۰)"],
        ["هسته", "sing-box — پشتیبان VLESS/REALITY، Hysteria2، TUIC، "
                 "Trojan، VMess و Shadowsocks"],
        ["کاربرد", "دسترسی به تلگرام، اینستاگرام، یوتیوب و کل وب"],
        ["سیستم عامل", "ویندوز ۱۰ و ۱۱"],
        ["پورت محلی", "127.0.0.1:10808 (SOCKS5 و HTTP)"],
    ], widths=[3.6, 12.4], ltr_cols=())

    callout(doc, "چهار نکته‌ی کلیدی",
            "۱) مرورگرها با فعال‌بودن «پروکسی سیستم» خودکار عبور می‌کنند. "
            "۲) برای تلگرام، دکمه‌ی «تلگرام» در صفحه‌ی اصلی را بزنید (بخش ۳). "
            "۳) برای عبور کل سیستم، حالت TUN را از تنظیمات فعال کنید. "
            "۴) اگر اشتراک شخصی دارید، در تنظیمات وارد کنید تا اول از همه تست شود.")

    doc.add_page_break()

    # ---------------- ۱. شروع سریع ----------------
    h1(doc, "شروع سریع در سه گام", 1)
    table(doc, ["گام", "کاری که انجام می‌دهید", "نتیجه"], [
        ["۱", "روی run.bat دوبار کلیک کنید (یا python main.py)",
         "پنجره‌ی مهراس باز می‌شود و هسته در اولین اجرا خودکار نصب می‌شود."],
        ["۲", "دکمه‌ی بزرگ وسط را بزنید", "برنامه گره‌ها را می‌گیرد، تست می‌کند و "
                                         "به بهترین وصل می‌شود."],
        ["۳", "از پنجره‌ی مرورگر خود استفاده کنید", "تلگرام/اینستاگرام/یوتیوب "
                                                   "باز می‌شوند."],
    ], widths=[1.4, 7.0, 7.6])

    par(doc, "اتصال اول معمولاً ۲۰ تا ۶۰ ثانیه طول می‌کشد: برنامه یک دور سریع "
             "می‌زند، به اولین گره‌ی سالم وصل می‌شود و بعد در پس‌زمینه گره‌های "
             "بیشتری را می‌سنجد تا اگر گره‌ی بهتری بود، خودش جابه‌جا شود. از اجرای "
             "دوم به بعد، اتصال از حافظه در چند ثانیه انجام می‌شود.")

    h2(doc, "رنگ‌ها و نشانه‌ها")
    bullet(doc, "یعنی اتصال برقرار و تونل سالم است.", bold_prefix="حلقه‌ی سبز: ")
    bullet(doc, "یعنی برنامه در حال گرفتن و تست گره‌ها است.",
           bold_prefix="حلقه‌ی زرد: ")
    bullet(doc, "یعنی چراغ سبز نشده؛ بخش ۶ (عیب‌یابی) را ببینید.",
           bold_prefix="نشان «بسته» روی یک سایت: ")
    bullet(doc, "ستون «تأخیر» پینگ واقعی تونل است؛ کمتر از ۹۰۰ میلی‌ثانیه خوب است.",
           bold_prefix="جدول سرورها: ")

    # ---------------- ۲. دو حالت اتصال ----------------
    h1(doc, "دو حالت اتصال: کدام را انتخاب کنم؟", 2)
    table(doc, ["حالت", "چه چیزی عبور می‌کند", "نیاز به دسترسی مدیر", "توصیه"], [
        ["پروکسی سیستم (پیش‌فرض)", "مرورگرها و برنامه‌هایی که پروکسی ویندوز را "
                                   "می‌خوانند", "خیر", "سریع‌ترین شروع"],
        ["TUN (کل سیستم)", "همه‌ی برنامه‌ها: تلگرام، دیسکورد، بازی، ترمینال",
         "بله (یک‌بار تأیید UAC)", "تجربه‌ی کامل وی‌پی‌ان"],
    ], widths=[3.6, 6.6, 3.4, 3.4])

    callout(doc, "چطور حالت TUN را فعال کنم؟",
            "تنظیمات ← حالت اتصال ← «حالت کامل TUN» ← ذخیره. سپس یک بار قطع و "
            "وصل کنید. اگر ویندوز اجازه‌ی مدیر بدهد، همه‌ی برنامه‌ها بدون تنظیم "
            "اضافه عبور می‌کنند. اگر اجازه نداده شود، برنامه به‌طور خودکار به "
            "حالت پروکسی برمی‌گردد و پیام می‌دهد.")

    # ---------------- ۳. تلگرام ----------------
    h1(doc, "اتصال تلگرام دسکتاپ", 3)
    par(doc, "تلگرام از پروکسی سیستم ویندوز پیروی نمی‌کند و باید یک بار دستی "
             "تنظیم شود. ساده‌ترین راه: در صفحه‌ی اصلی مهراس دکمه‌ی «تلگرام» را "
             "بزنید؛ آدرس و پورت آماده است و با یک کلیک کپی می‌شود. اگر ترجیح "
             "می‌دهید دستی انجام دهید:")
    table(doc, ["گام", "مسیر"], [
        ["۱", "در تلگرام: Settings ← Advanced ← Connection type"],
        ["۲", "گزینه‌ی «TCP with custom proxy» یا «Use custom proxy» را انتخاب کنید"],
        ["۳", "نوع پروکسی را SOCKS5 بگذارید"],
        ["۴", "Host: 127.0.0.1    Port: 10808"],
        ["۵", "Save کنید؛ کنار نام شما باید علامت اتصال ظاهر شود"],
    ], widths=[1.4, 14.6], ltr_cols=(1,))

    callout(doc, "نکته",
            "اگر پورت را در تنظیمات مهراس تغییر دادید، همان عدد را در تلگرام هم "
            "بگذارید. برای تلگرام موبایل، اگر گوشی و کامپیوتر در یک شبکه هستند، "
            "می‌توانید از حالت TUN استفاده کنید یا در گوشی یک کلاینت استاندارد "
            "با همان لینک اشتراک بسازید.")

    # ---------------- ۴. مرورگر ----------------
    h1(doc, "اینستاگرام و یوتیوب", 4)
    par(doc, "این دو سایت از طریق مرورگر باز می‌شوند و به تنظیم جداگانه نیاز "
             "ندارند؛ کافی است در مهراس گزینه‌ی «اعمال خودکار پروکسی روی ویندوز» "
             "روشن باشد (پیش‌فرض).")
    bullet(doc, "اگر سایت باز نشد، اول از سالم بودن تونل مطمئن شوید: کارت "
                "«دسترسی سایت‌ها» باید سه نشان سبز داشته باشد.")
    bullet(doc, "برنامه‌های نصب‌شده (اپ اینستاگرام ویندوز، اپ یوتیوب) فقط در "
                "حالت TUN عبور می‌کنند.")
    bullet(doc, "اگر مرورگر دیگری پروکسی مخصوص خود دارد (مثل Firefox)، آن را روی "
                "HTTP Proxy با آدرس 127.0.0.1 و پورت 10808 تنظیم کنید.")

    # ---------------- ۵. سرورها و تنظیمات ----------------
    h1(doc, "انتخاب سرور و تنظیمات مفید", 5)
    par(doc, "برنامه به‌طور خودکار بهترین گره را انتخاب می‌کند، اما دست شما هم "
             "باز است:")
    bullet(doc, "برای اتصال به یک سرور خاص، روی همان سطر دوبار کلیک کنید.")
    bullet(doc, "ستون «سرعت» فقط برای گره‌های برتر اندازه‌گیری می‌شود؛ خالی بودن "
                "آن به معنی خرابی نیست.")
    bullet(doc, "با «↻ اسکن گره‌های تازه» یک دور تازه‌ی دریافت و تست شروع می‌شود.")
    bullet(doc, "گزینه‌ی «جابه‌جایی خودکار به گره‌ی بهتر» باعث می‌شود اگر گره‌ی "
                "بهتری پیدا شد، برنامه خودش جابه‌جا شود.")

    h2(doc, "اشتراک یا کانفیگ شخصی (مهم‌ترین توصیه)")
    par(doc, "اگر یک اشتراک خریداری‌شده یا کانفیگ شخصی دارید، آن را در "
             "«تنظیمات ← اشتراک یا کانفیگ شخصی» وارد کنید: هر خط یک آدرس اشتراک "
             "(https) یا یک لینک کانفیگ (vless، trojan، ss، hysteria2 و…). "
             "این گره‌ها همیشه اول تست می‌شوند و امتیاز بالاتری می‌گیرند، چون "
             "معمولاً از گره‌های عمومی رایگان پایدارترند. برای بارگذاری فوری، "
             "بعد از ذخیره‌ی تنظیمات یک اسکن تازه شروع می‌شود.")

    h2(doc, "تنظیمات پیشنهادی")
    table(doc, ["تنظیم", "پیش‌فرض", "توضیح"], [
        ["حالت اتصال", "پروکسی سیستم", "برای شروع؛ TUN برای عبور کل سیستم"],
        ["اعمال پروکسی ویندوز", "روشن", "مرورگرها را خودکار عبور می‌دهد"],
        ["سایت‌های ایرانی مستقیم", "روشن", "سایت‌های داخلی سریع و بدون مصرف حجم"],
        ["فاصله‌ی بررسی سلامت", "۲۰ ثانیه", "اگر اتصال بمیرد، سریع جایگزین می‌شود"],
        ["فاصله‌ی اسکن دوره‌ای", "۶۰ دقیقه", "گره‌های تازه‌تر پیدا می‌شوند"],
        ["حداکثر گره برای تست", "۱۸۰", "بیشتر = دقیق‌تر ولی کندتر"],
    ], widths=[4.4, 3.0, 9.6])

    # ---------------- ۶. عیب‌یابی ----------------
    h1(doc, "عیب‌یابی سریع", 6)
    table(doc, ["نشانه", "علت احتمالی", "راه‌حل"], [
        ["«هیچ گره‌ی سالمی پیدا نشد»", "اینترنت قطع است یا منابع موقتاً در دسترس "
                                       "نیستند", "اتصال اینترنت را بررسی کنید و "
                                                 "چند دقیقه بعد دوباره اسکن کنید."],
        ["اتصال برقرار است ولی سایت باز نمی‌شود", "پروکسی سیستم اعمال نشده یا "
                                                  "مرورگر کش کرده", "تنظیمات ← "
                                                                   "اعمال پروکسی روشن؛ مرورگر را ری‌استارت کنید."],
        ["سرعت پایین است", "گره‌ی دور یا شلوغ", "اسکن تازه بزنید؛ گره‌ی با تأخیر "
                                               "کمتر و سرعت بالاتر انتخاب می‌شود."],
        ["هر چند دقیقه قطع می‌شود", "گره ناپایدار است", "«جابه‌جایی خودکار» را "
                                                       "روشن بگذارید؛ برنامه خودش جایگزین می‌کند."],
        ["حالت TUN فعال نمی‌شود", "ویندوز اجازه‌ی مدیر نداده", "با «Run as "
                                                              "administrator» اجرا کنید یا در حالت پروکسی بمانید."],
        ["تلگرام وصل نمی‌شود", "پروکسی تلگرام تنظیم نشده", "بخش ۳ را انجام دهید "
                                                          "(SOCKS5 روی 127.0.0.1:10808)."],
        ["سایت بانکی باز نمی‌شود", "IP خارجی برای بانک مشکل‌ساز است", "«سایت‌های "
                                                                     "ایرانی مستقیم» را روشن کنید."],
    ], widths=[4.6, 5.0, 7.4], font_size=9)

    # ---------------- ۷. امنیت ----------------
    h1(doc, "امنیت، حریم خصوصی و مصرف", 7)
    bullet(doc, "کانفیگ‌ها از مخازن عمومی گرفته می‌شوند؛ برای اطلاعات حساس "
                "(بانک، ایمیل کاری) روی گره‌های عمومی حساب ویژه باز نکنید.")
    bullet(doc, "مهراس هیچ داده‌ای از شما را به سرور خودش نمی‌فرستد؛ آمار مصرف "
                "فقط روی همین کامپیوتر ذخیره می‌شود.")
    bullet(doc, "با روشن بودن «سایت‌های ایرانی مستقیم»، ترافیک داخلی از تونل "
                "عبور نمی‌کند و حجم کمتری مصرف می‌شود.")
    bullet(doc, "تنظیمات پروکسی قبلی ویندوز ذخیره و هنگام قطع به حالت اول "
                "برگردانده می‌شود.")
    bullet(doc, "گره‌های عمومی رایگان‌اند و ممکن است چند ساعت یا چند روز عمر "
                "کنند؛ به همین دلیل اسکن دوره‌ای و جایگزینی خودکار وجود دارد.")

    # ---------------- ۸. پرسش‌های متداول ----------------
    doc.add_page_break()
    h1(doc, "پرسش‌های متداول", 8)
    faq = [
        ("آیا باید کانفیگ وارد کنم؟",
         "نه. برنامه منابع را خودش می‌گیرد و تست می‌کند."),
        ("چرا بار اول کند است؟",
         "هزاران گره دریافت و صدها گره با اتصال واقعی آزمایش می‌شوند. از اجرای "
         "بعد، اتصال از حافظه انجام می‌شود."),
        ("چرا بعضی سرورها «ناموفق» هستند؟",
         "بیشتر گره‌های عمومی قدیمی یا مسدود شده‌اند؛ برنامه آن‌ها را نشان می‌دهد "
         "ولی هرگز به‌طور خودکار انتخاب نمی‌کند."),
        ("فایل‌های برنامه کجا هستند؟",
         "data/logs/app.log (رویدادها)، data/logs/singbox.log (خروجی هسته) و "
         "data/working_cache.json (گره‌های سالم)."),
        ("آیا روی چند دستگاه کار می‌کند؟",
         "پوشه‌ی برنامه را روی دستگاه دیگر هم اجرا کنید؛ هر دستگاه مستقل کار "
         "می‌کند."),
    ]
    for q, a in faq:
        rich_par(doc, [(f"پرسش: {q}", {"bold": True, "color": ACCENT})],
                 align="right", space_after=1)
        par(doc, f"پاسخ: {a}", space_after=7)

    # ---------------- اطلاعات فنی ----------------
    h1(doc, "اطلاعات فنی برای علاقه‌مندان", 9)
    table(doc, ["مورد", "مقدار"], [
        ["پورت SOCKS5/HTTP", "127.0.0.1:10808"],
        ["پورت HTTP جداگانه", "127.0.0.1:10809"],
        ["هسته", "sing-box (خودکار نصب می‌شود)"],
        ["پروتکل‌های پشتیبانی‌شده", "VLESS + REALITY، Hysteria2، TUIC، AnyTLS، "
                                     "Trojan، VMess، Shadowsocks"],
        ["فایل تنظیمات برنامه", "data/settings.json"],
        ["گره‌های سالم ذخیره‌شده", "data/working_cache.json"],
        ["لاگ‌ها", "data/logs"],
        ["رابط کاربری", "PySide6 (Qt 6) — راست‌به‌چپ"],
    ], widths=[4.6, 11.4])

    par(doc, "", space_after=10)
    end = par(doc, "", align="center")
    _style_run(end.add_run("مهراس — اینترنت آزاد، بدون پیچیدگی"), size=10,
               color=GRAY)

    # --- پاک‌سازی فراداده‌ها: هیچ نام/ایمیل/مسیر شخصی در فایل منتشرشده نباشد ---
    props = doc.core_properties
    props.author = "Mehras VPN"
    props.last_modified_by = "Mehras VPN"
    props.title = "راهنمای استفاده — وی پی ان شخصی مهراس"
    props.subject = "Mehras VPN"
    props.keywords = "VPN, sing-box, Windows, راهنما"
    props.comments = ""
    props.category = "User guide"
    props.identifier = "mehras-vpn-guide"
    props.revision = 1

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print("saved:", OUT)


if __name__ == "__main__":
    build()
